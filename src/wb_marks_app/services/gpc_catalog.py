from __future__ import annotations

from dataclasses import dataclass
from importlib import resources
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from wb_marks_app.server_models import GpcCatalogModel


@dataclass(frozen=True)
class GpcCatalogItem:
    code: str
    description: str

    def to_dict(self) -> dict[str, str]:
        return {
            "code": self.code,
            "description": self.description,
            "label": f"{self.code} — {self.description}",
        }


def seed_gpc_catalog(session: Session) -> int:
    if not hasattr(session, "execute"):
        return 0

    existing_count = session.scalar(select(func.count()).select_from(GpcCatalogModel)) or 0
    if existing_count > 0:
        return 0

    items = _asset_gpc_catalog_items()
    for item in items:
        session.add(
            GpcCatalogModel(
                code=item.code,
                description=item.description,
                source_file="gpc_catalog.xlsx",
            )
        )
    return len(items)


def search_gpc_catalog(session: Session, query: str = "", limit: int = 50) -> list[dict[str, str]]:
    if not hasattr(session, "execute"):
        return []

    safe_limit = max(1, min(int(limit or 50), 100))
    query_text = _text(query).casefold()
    terms = [term for term in query_text.split() if term]
    result: list[GpcCatalogItem] = []

    rows = session.execute(select(GpcCatalogModel).order_by(GpcCatalogModel.code)).scalars()
    for row in rows:
        item = GpcCatalogItem(code=row.code, description=row.description)
        haystack = f"{item.code} {item.description}".casefold()
        if not terms or all(term in haystack for term in terms):
            result.append(item)
        if len(result) >= safe_limit:
            break

    return [item.to_dict() for item in result]


def gpc_catalog_description(session: Session, code: str) -> str:
    if not hasattr(session, "get"):
        return ""
    normalized = _normalize_code(code)
    if not normalized:
        return ""
    row = session.get(GpcCatalogModel, normalized)
    return row.description if row else ""


def _asset_gpc_catalog_items() -> tuple[GpcCatalogItem, ...]:
    catalog_path = resources.files("wb_marks_app").joinpath("assets/gpc_catalog.xlsx")
    with resources.as_file(catalog_path) as path:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        items: list[GpcCatalogItem] = []
        for row in worksheet.iter_rows(min_row=3, values_only=True):
            code = _normalize_code(row[0] if len(row) > 0 else "")
            description = _text(row[1] if len(row) > 1 else "")
            if code and description:
                items.append(GpcCatalogItem(code=code, description=description))
        workbook.close()
    return tuple(items)


def _normalize_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    code = "".join(ch for ch in str(value).strip() if ch.isdigit())
    return code if len(code) == 8 else ""


def _text(value: Any) -> str:
    return str(value or "").strip()
