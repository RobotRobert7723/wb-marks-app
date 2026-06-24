from __future__ import annotations

import warnings
from copy import copy
from io import BytesIO
from importlib import resources
from typing import Any

from openpyxl import load_workbook


DATA_SHEET_NAME = "1. данные_о_продукте"
DATA_START_ROW = 7
DATA_COLUMNS = 22


def build_gtin_request_workbook(product_card: Any, gpc_code: str) -> bytes:
    template_path = resources.files("wb_marks_app").joinpath("assets/gtin_request_template.xlsx")
    with resources.as_file(template_path) as path:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
            workbook = load_workbook(path)

    worksheet = workbook[DATA_SHEET_NAME]
    rows = list(product_card.rows)
    _prepare_data_rows(worksheet, max(len(rows), 1))

    brand = _text(product_card.wb_summary.brand)
    functional_name = _capitalize_text(product_card.wb_summary.seller_category or product_card.wb_summary.name)
    seller_article = _text(product_card.wb_summary.seller_article)
    color = _text(product_card.wb_summary.color).lower()

    if not rows:
        rows = []
    for offset, row in enumerate(rows):
        excel_row = DATA_START_ROW + offset
        size = _text(row.wb_size)
        row_color = _text(row.color).lower() or color
        values = [
            brand,
            None,
            "Русский",
            functional_name,
            _variant_text(seller_article, row_color, size),
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "Россия (Российская Федерация)",
            "Кыргызстан",
            None,
            _text(gpc_code),
            "1",
            "Штуки",
            "Потребительский продукт",
            "Фиксированное измерение",
        ]
        for col_index, value in enumerate(values, start=1):
            worksheet.cell(row=excel_row, column=col_index, value=value)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _prepare_data_rows(worksheet, row_count: int) -> None:
    row_count = max(row_count, 1)
    if worksheet.max_row > DATA_START_ROW:
        worksheet.delete_rows(DATA_START_ROW + 1, worksheet.max_row - DATA_START_ROW)
    if row_count > 1:
        worksheet.insert_rows(DATA_START_ROW + 1, row_count - 1)

    template_cells = [worksheet.cell(DATA_START_ROW, column) for column in range(1, DATA_COLUMNS + 1)]
    template_height = worksheet.row_dimensions[DATA_START_ROW].height
    for row in range(DATA_START_ROW, DATA_START_ROW + row_count):
        if template_height is not None:
            worksheet.row_dimensions[row].height = template_height
        for column, template_cell in enumerate(template_cells, start=1):
            cell = worksheet.cell(row=row, column=column)
            if template_cell.has_style:
                cell._style = copy(template_cell._style)
            if template_cell.number_format:
                cell.number_format = template_cell.number_format
            if template_cell.alignment:
                cell.alignment = copy(template_cell.alignment)
            if template_cell.protection:
                cell.protection = copy(template_cell.protection)


def _variant_text(seller_article: str, color: str, size: str) -> str:
    parts = []
    if seller_article:
        parts.append(f"Арт.{seller_article}")
    if color:
        parts.append(f"цвет: {color}")
    if size:
        parts.append(f"р. {size}")
    return ", ".join(parts)


def _capitalize_text(value: str) -> str:
    text = _text(value)
    if not text:
        return ""
    lowered = text.lower()
    return lowered[:1].upper() + lowered[1:]


def _text(value: Any) -> str:
    return str(value or "").strip()
