from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import BinaryIO

from wb_marks_app.exceptions import IntegrationUnavailableError


@dataclass(frozen=True, slots=True)
class GtinExcelRow:
    gtin: str
    brand: str
    functional_name: str
    variety: str
    vendor_article: str
    color: str
    size: str


class GtinExcelParser:
    def parse(self, file: bytes | BinaryIO) -> list[GtinExcelRow]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise IntegrationUnavailableError("The openpyxl package is not installed.") from exc

        source = BytesIO(file) if isinstance(file, bytes) else file
        try:
            workbook = load_workbook(source, data_only=True, read_only=True)
        except Exception as exc:
            raise ValueError("Не удалось прочитать Excel файл.") from exc

        rows: list[GtinExcelRow] = []
        for sheet in workbook.worksheets:
            header_row_index, columns = self._find_header(sheet)
            if not columns:
                continue
            for values in sheet.iter_rows(min_row=header_row_index + 1, values_only=True):
                row = self._row_from_values(values, columns)
                if row is not None:
                    rows.append(row)
        return rows

    def find_by_vendor_article(self, rows: list[GtinExcelRow], vendor_article: str) -> list[GtinExcelRow]:
        normalized_article = self.normalize_article(vendor_article)
        if not normalized_article:
            return []
        return [row for row in rows if self.normalize_article(row.vendor_article) == normalized_article]

    def normalize_article(self, value: str) -> str:
        return re.sub(r"\s+", "", str(value or "").strip()).casefold()

    def _find_header(self, sheet) -> tuple[int, dict[str, int]]:
        for row_index, values in enumerate(sheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1):
            columns: dict[str, int] = {}
            for index, value in enumerate(values):
                header = self._normalize_header(value)
                if not header:
                    continue
                if header == "GTIN":
                    columns.setdefault("gtin", index)
                elif "СУБ-БРЕНД" not in header and header.endswith("БРЕНД"):
                    columns.setdefault("brand", index)
                elif self._is_primary_header(header, "ФУНКЦИОНАЛЬНОЕ НАЗВАНИЕ"):
                    columns.setdefault("functional_name", index)
                elif self._is_primary_header(header, "РАЗНОВИДНОСТЬ"):
                    columns.setdefault("variety", index)
            if {"gtin", "brand", "functional_name", "variety"}.issubset(columns):
                return row_index, columns
        return 0, {}

    def _row_from_values(self, values: tuple, columns: dict[str, int]) -> GtinExcelRow | None:
        gtin = self._cell_text(values, columns["gtin"])
        brand = self._cell_text(values, columns["brand"])
        functional_name = self._cell_text(values, columns["functional_name"])
        variety = self._cell_text(values, columns["variety"])
        if not gtin and not brand and not functional_name and not variety:
            return None

        parsed = self._parse_variety(variety)
        return GtinExcelRow(
            gtin=gtin,
            brand=brand,
            functional_name=functional_name,
            variety=variety,
            vendor_article=parsed["vendor_article"],
            color=parsed["color"],
            size=parsed["size"],
        )

    def _parse_variety(self, value: str) -> dict[str, str]:
        article = self._match_group(value, r"(?:^|[,;]\s*)арт\.?\s*([^,;]+)")
        color = self._match_group(value, r"(?:^|[,;]\s*)цвет\s*[:\-]?\s*([^,;]+)")
        size = self._match_group(value, r"(?:^|[,;]\s*)(?:р\.?|размер)\s*[:\-]?\s*([^,;]+)")
        return {
            "vendor_article": article,
            "color": color,
            "size": size,
        }

    def _match_group(self, value: str, pattern: str) -> str:
        match = re.search(pattern, value or "", flags=re.IGNORECASE)
        return self._clean_value(match.group(1)) if match else ""

    def _is_primary_header(self, header: str, expected: str) -> bool:
        if expected not in header:
            return False
        return not re.search(r"(^|\s|-)2(\s|$)|(^|\s|-)3(\s|$)", header)

    def _normalize_header(self, value) -> str:
        text = str(value or "").replace("Ё", "Е").upper()
        return re.sub(r"\s+", " ", text).strip()

    def _cell_text(self, values: tuple, index: int) -> str:
        if index >= len(values):
            return ""
        return self._clean_value(values[index])

    def _clean_value(self, value) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return re.sub(r"\s+", " ", str(value).strip())
