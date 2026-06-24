import json
import unittest
from dataclasses import replace
from datetime import date, datetime, timezone
from base64 import b64encode
from io import BytesIO
from tempfile import TemporaryDirectory

from fastapi.testclient import TestClient
from itsdangerous import TimestampSigner
from openpyxl import Workbook


class ProductCardRouteTests(unittest.TestCase):
    def test_product_card_page_uses_wb_article_from_url(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_teksher_mapping_service = webapp.teksher_mapping_service
        old_teksher_product_service = webapp.teksher_product_service
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.teksher_mapping_service = FakeTeksherMappingService()
        webapp.teksher_product_service = FakeTeksherProductService()
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="Данные WB загружены из Content API",
                wb_summary=WbProductSummary(
                    name="Спортивный костюм",
                    seller_category="Костюмы спортивные",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="Кыргызстан",
                    seller_article="cv_nk_white_smr",
                    color="белый",
                    composition="полиэстер 100%",
                    gender="мальчики",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462689",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="38 МЕЖДУНАРОДНЫЙ",
                        product_type="КОСТЮМЫ СПОРТИВНЫЕ",
                        gtin="",
                        tnved="6112120000",
                        country="Кыргызстан",
                        vendor_article="cv_nk_white_smr",
                        color="БЕЛЫЙ",
                        composition="полиэстер 100%",
                        target_gender="МАЛЬЧИКИ",
                        trademark="ErLine",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=7,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            page = client.get("/product-cards/847012873")
            self.assertEqual(200, page.status_code)
            self.assertIn("Карточка WB 847012873", page.text)
            self.assertIn('data-wb-article="847012873"', page.text)
            self.assertIn("/product-cards/847012873", page.text)
            self.assertIn(".mapping-table .ready-print-head { text-align: left; }", page.text)
            self.assertIn(".mapping-table .ready-print-cell { text-align: left; }", page.text)
            self.assertIn('class="template-head ready-print-head"', page.text)
            self.assertIn("Готовы к печати", page.text)
            self.assertNotIn("Готовы к нанесению", page.text)
            self.assertNotIn("<th class=\"template-head\">Напечатать</th>", page.text)
            self.assertIn('class="template-cell ready-print-cell" data-ready-print', page.text)
            self.assertIn("data-ready-print", page.text)
            self.assertIn("Заказ ЧЗ в Текшер", page.text)
            self.assertIn('data-order-gtin', page.text)
            self.assertIn('class="order-count-input"', page.text)
            self.assertIn('value="0"', page.text)
            self.assertNotIn('value="7"', page.text)
            self.assertIn("Трансгран", page.text)
            self.assertIn('class="transgran-checkbox"', page.text)
            self.assertIn("checked", page.text)
            self.assertIn("Статус", page.text)
            self.assertIn('id="mark-order-button"', page.text)
            self.assertIn("data-order-status", page.text)
            self.assertIn("mark-orders", page.text)
            self.assertIn('id="order-history-button"', page.text)
            self.assertIn('id="order-history-panel" class="card order-history-panel"', page.text)
            self.assertIn("font-size: 1.5em;", page.text)
            self.assertNotIn("<th class=\"template-head\">Документ</th>", page.text)
            self.assertNotIn("Последние операции сверху", page.text)
            self.assertNotIn("mark-order-initial-state", page.text)
            self.assertNotIn("hydrateInitialMarkOrderState", page.text)
            self.assertIn('data-gtin-field="gtin"', page.text)
            self.assertIn('data-gtin-field="product_type"', page.text)
            self.assertIn('data-gtin-field="full_name"', page.text)
            self.assertIn('data-full-name=', page.text)
            self.assertIn('<th class="teksher-head">Вид товара</th>', page.text)
            self.assertIn('<th class="teksher-head">Полное наименование</th>', page.text)
            self.assertIn('data-gtin-field="tnved"', page.text)
            self.assertIn('data-gtin-field="target_gender"', page.text)
            self.assertIn('data-wb-product-type="Костюмы спортивные"', page.text)
            self.assertIn('data-wb-tnved="6112120000"', page.text)
            self.assertIn('data-wb-country="Кыргызстан"', page.text)
            self.assertIn('data-wb-color="белый"', page.text)
            self.assertIn('data-wb-composition="полиэстер 100%"', page.text)
            self.assertIn('data-wb-brand="ErLine"', page.text)
            self.assertIn('data-has-mapping="false"', page.text)
            self.assertIn('id="mapping-save-button" class="secondary" disabled', page.text)
            self.assertIn('id="label-template-select"', page.text)
            self.assertIn('<option value="srad" selected>SRad</option>', page.text)
            self.assertIn('id="print-labels-button"', page.text)
            self.assertIn('id="print-history-button"', page.text)
            self.assertIn('id="print-history-panel" class="card order-history-panel print-history-panel"', page.text)
            self.assertIn("/label-print", page.text)
            self.assertIn("/label-prints", page.text)
            self.assertIn("Для этой карточки уже есть GTIN. Действительно хотите обновить", page.text)
            self.assertIn("В таблице есть красные поля. Сохранить?", page.text)
            self.assertIn("gtin-loaded", page.text)
            self.assertIn("gtin-mismatch", page.text)
            self.assertIn("function setLoadedCell(", page.text)
            self.assertIn("function clearGtinTable()", page.text)
            self.assertIn("function targetGenderFromWb(", page.text)

            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.teksher_mapping_service = old_teksher_mapping_service
            webapp.teksher_product_service = old_teksher_product_service

    def test_product_card_page_uses_teksher_data_for_saved_gtin(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_teksher_mapping_service = webapp.teksher_mapping_service
        old_teksher_product_service = webapp.teksher_product_service
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.teksher_mapping_service = FakeTeksherMappingService(
            version=3,
            latest_rows=[
                {
                    "wb_barcode": "2049271462689",
                    "wb_size": "38",
                    "wb_ru_size": "134",
                    "teksher_size": "STALE SIZE",
                    "product_type": "STALE PRODUCT",
                    "gtin": "04709055620626",
                    "tnved": "STALE TNVED",
                    "country": "STALE COUNTRY",
                    "vendor_article": "STALE ARTICLE",
                    "color": "STALE COLOR",
                    "composition": "STALE COMPOSITION",
                    "target_gender": "STALE GENDER",
                    "trademark": "STALE BRAND",
                }
            ],
        )
        webapp.teksher_product_service = FakeTeksherProductService(
            mapping_rows_by_gtin={
                "04709055620626": {
                    "teksher_size": "38 МЕЖДУНАРОДНЫЙ",
                    "product_type": "ДАННЫЕ ИЗ ТЕКШЕР",
                    "gtin": "04709055620626",
                    "tnved": "6112120000",
                    "country": "Кыргызстан",
                    "vendor_article": "cv_nk_white_smr",
                    "color": "БЕЛЫЙ",
                    "composition": "полиэстер 100%",
                    "target_gender": "МУЖСКОЙ",
                    "trademark": "ErLine",
                }
            }
        )
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Спортивный костюм",
                    seller_category="Костюмы спортивные",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="Кыргызстан",
                    seller_article="cv_nk_white_smr",
                    color="белый",
                    composition="полиэстер 100%",
                    gender="мальчики",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462689",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            page = client.get("/product-cards/847012873")

            self.assertEqual(200, page.status_code)
            self.assertIn("ДАННЫЕ ИЗ ТЕКШЕР", page.text)
            self.assertIn("04709055620626", page.text)
            self.assertNotIn("STALE PRODUCT", page.text)
            self.assertEqual(["04709055620626"], webapp.teksher_product_service.gtins)
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.teksher_mapping_service = old_teksher_mapping_service
            webapp.teksher_product_service = old_teksher_product_service

    def test_gtin_upload_returns_rows_for_matching_vendor_article(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="",
                    seller_category="",
                    wb_article="847012873",
                    tnved="",
                    country="",
                    seller_article="cv_nk_white_smr",
                    color="",
                    composition="",
                    gender="",
                    brand="",
                ),
                rows=[],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/gtin-upload",
                files={
                    "file": (
                        "gtin.xlsx",
                        _build_gtin_upload_workbook("Арт.cv_nk_white_smr, цвет: белый, р. 38"),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual("04709055620626", payload["rows"][0]["gtin"])
            self.assertEqual("cv_nk_white_smr", payload["rows"][0]["vendor_article"])
            self.assertEqual("38", payload["rows"][0]["size"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service

    def test_gtin_upload_reports_missing_vendor_article(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="",
                    seller_category="",
                    wb_article="847012873",
                    tnved="",
                    country="",
                    seller_article="cv_nk_white_smr",
                    color="",
                    composition="",
                    gender="",
                    brand="",
                ),
                rows=[],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/gtin-upload",
                files={
                    "file": (
                        "gtin.xlsx",
                        _build_gtin_upload_workbook("Арт.other_article, цвет: белый, р. 38"),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertFalse(payload["ok"])
            self.assertEqual("Артикул продавца в файле не найден", payload["message"])
            self.assertEqual([], payload["rows"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service

    def test_mapping_preview_endpoint_returns_draft_dialog_payload(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_teksher_product_service = webapp.teksher_product_service
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.teksher_product_service = FakeTeksherProductService(existing_gtins={"04709055620664"})
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Sport suit",
                    seller_category="Sport suits",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="KG",
                    seller_article="cv_nk_white_smr",
                    color="white",
                    composition="polyester 100%",
                    gender="boys",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462689",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/mapping/preview",
                json={
                    "source": "gtin_excel",
                    "rows": [
                        {"wb_size": "38", "gtin": "04709055620664"},
                        {"wb_size": "40", "gtin": "04709055620671"},
                    ],
                },
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual(["04709055620664"], payload["existing_gtins"])
            self.assertEqual(["04709055620671"], payload["create_gtins"])
            self.assertEqual("КЫРГЫЗСТАН", payload["draft_fields"]["country"])
            self.assertIn({"value": "МУЖСКОЙ", "label": ""}, payload["dictionaries"]["target_gender"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.teksher_product_service = old_teksher_product_service

    def test_mapping_save_endpoint_creates_version(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_teksher_mapping_service = webapp.teksher_mapping_service
        old_teksher_product_service = webapp.teksher_product_service
        fake_mapping_service = FakeTeksherMappingService(version=2)
        fake_teksher_product_service = FakeTeksherProductService(draft_ids=["draft-1"])
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.teksher_mapping_service = fake_mapping_service
        webapp.teksher_product_service = fake_teksher_product_service
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Sport suit",
                    seller_category="Sport suits",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="KG",
                    seller_article="cv_nk_white_smr",
                    color="white",
                    composition="polyester 100%",
                    gender="boys",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462689",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/mapping",
                json={
                    "source": "gtin_excel",
                    "rows": [
                        {
                            "wb_barcode": "2049271462689",
                            "wb_size": "38",
                            "wb_ru_size": "134",
                            "teksher_size": "38 МЕЖДУНАРОДНЫЙ",
                            "product_type": "КОСТЮМ СПОРТИВНЫЙ",
                            "gtin": "04709055620626",
                            "vendor_article": "cv_nk_white_smr",
                            "color": "БЕЛЫЙ",
                            "trademark": "ErLine",
                        }
                    ],
                },
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual(2, payload["version"])
            self.assertEqual(1, payload["rows_saved"])
            self.assertEqual("user-1", fake_mapping_service.saved[0]["user_id"])
            self.assertEqual("gtin_excel", fake_mapping_service.saved[0]["source"])
            self.assertEqual("04709055620626", fake_mapping_service.saved[0]["rows"][0]["gtin"])
            self.assertEqual(["draft-1"], payload["teksher_draft_ids"])
            self.assertEqual("04709055620626", fake_teksher_product_service.calls[0]["rows"][0]["gtin"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.teksher_mapping_service = old_teksher_mapping_service
            webapp.teksher_product_service = old_teksher_product_service

    def test_mapping_save_endpoint_saves_when_teksher_gtin_exists(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_teksher_mapping_service = webapp.teksher_mapping_service
        old_teksher_product_service = webapp.teksher_product_service
        fake_mapping_service = FakeTeksherMappingService(version=2)
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.teksher_mapping_service = fake_mapping_service
        webapp.teksher_product_service = FakeTeksherProductService(existing=True)
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Sport suit",
                    seller_category="Sport suits",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="KG",
                    seller_article="cv_nk_white_smr",
                    color="white",
                    composition="polyester 100%",
                    gender="boys",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462689",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/mapping",
                json={
                    "source": "gtin_excel",
                    "rows": [
                        {
                            "wb_barcode": "2049271462689",
                            "wb_size": "38",
                            "wb_ru_size": "134",
                            "gtin": "04709055620626",
                        }
                    ],
                },
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual(2, payload["version"])
            self.assertEqual(1, payload["rows_saved"])
            self.assertEqual([], payload["teksher_draft_ids"])
            self.assertEqual(1, len(fake_mapping_service.saved))
            self.assertEqual("04709055620626", fake_mapping_service.saved[0]["rows"][0]["gtin"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.teksher_mapping_service = old_teksher_mapping_service
            webapp.teksher_product_service = old_teksher_product_service

    def test_mark_order_endpoint_starts_product_card_run(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_workflow_service = webapp.workflow_service
        old_get_user_run = webapp._get_user_run
        old_serialize_product_card_order_run = webapp._serialize_product_card_order_run
        fake_workflow_service = FakeWorkflowRunService()
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.workflow_service = fake_workflow_service
        webapp._get_user_run = lambda _session, run_id, _user_id: type("Run", (), {"id": run_id, "status": "created"})()
        webapp._serialize_product_card_order_run = lambda run, _session: {
            "run": {"id": run.id, "status": run.status},
            "items": [{"size": "38", "gtin": "04709055620626", "status_text": "Ожидает запуска"}],
        }
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Sport suit",
                    seller_category="Sport suits",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="KG",
                    seller_article="cv_nk_white_smr",
                    color="white",
                    composition="polyester 100%",
                    gender="boys",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462689",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="04709055620626",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/mark-orders",
                json={
                    "rows": [
                        {
                            "size": "38",
                            "gtin": "04709055620626",
                            "quantity": 2,
                            "transgran": False,
                        }
                    ]
                },
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual("run-1", payload["run_id"])
            self.assertEqual("user-1", fake_workflow_service.calls[0]["user_id"])
            self.assertEqual("847012873", fake_workflow_service.calls[0]["wb_article"])
            self.assertEqual(2, fake_workflow_service.calls[0]["rows"][0]["quantity"])
            self.assertFalse(fake_workflow_service.calls[0]["rows"][0]["transgran"])
            self.assertEqual("Ожидает запуска", payload["status"]["items"][0]["status_text"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.workflow_service = old_workflow_service
            webapp._get_user_run = old_get_user_run
            webapp._serialize_product_card_order_run = old_serialize_product_card_order_run

    def test_mark_order_history_endpoint_returns_latest_first(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_product_card_order_history = webapp._product_card_order_history
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp._product_card_order_history = lambda _session, user_id, wb_article: [
            {"run": {"id": "new-run", "created_at": "2026-06-23T10:00:00"}, "items": []},
            {"run": {"id": "old-run", "created_at": "2026-06-23T09:00:00"}, "items": []},
        ]
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.get("/api/product-cards/847012873/mark-orders")

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual(["new-run", "old-run"], [item["run"]["id"] for item in payload["history"]])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp._product_card_order_history = old_product_card_order_history

    def test_label_print_endpoint_uses_product_card_pdf_helper(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_create_product_card_label_pdf = webapp._create_product_card_label_pdf
        calls = []
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp._create_product_card_label_pdf = lambda request, session, user_id, product_card, template: calls.append(
            {
                "user_id": user_id,
                "product_card": product_card,
                "template": template,
            }
        ) or {
            "ok": True,
            "download_url": "http://testserver/api/labels/pdf/file",
            "labels_count": 2,
            "history": [],
        }
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Sport suit",
                    seller_category="Sport suits",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="KG",
                    seller_article="cv_nk_blue_smr",
                    color="blue",
                    composition="polyester 100%",
                    gender="boys",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462634",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="04709055620664",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=1,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post("/api/product-cards/847012873/label-print", json={"template": "Medium"})

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual(2, payload["labels_count"])
            self.assertEqual("medium", calls[0]["template"])
            self.assertEqual("847012873", calls[0]["product_card"].wb_article)
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp._create_product_card_label_pdf = old_create_product_card_label_pdf

    def test_label_print_history_endpoint_returns_latest_first(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_product_card_label_print_history = webapp._product_card_label_print_history
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp._product_card_label_print_history = lambda request, _session, user_id, wb_article: [
            {"id": "new-print", "created_at": "2026-06-23T10:00:00", "template": "SRad"},
            {"id": "old-print", "created_at": "2026-06-23T09:00:00", "template": "Simple"},
        ]
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.get("/api/product-cards/847012873/label-prints")

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual(["new-print", "old-print"], [item["id"] for item in payload["history"]])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp._product_card_label_print_history = old_product_card_label_print_history

    def test_label_print_repeat_endpoint_uses_print_history_row(self) -> None:
        import wb_marks_app.webapp as webapp
        from wb_marks_app.models import AppConfig
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        old_create_all = webapp.create_all
        old_load_config = webapp.load_config
        old_session_scope = webapp.session_scope
        old_get_or_create_settings = webapp.get_or_create_settings
        old_settings_to_app_config = webapp.settings_to_app_config
        old_product_card_service = webapp.product_card_service
        old_teksher_mapping_service = webapp.teksher_mapping_service
        old_get_user_label_print_job = webapp._get_user_label_print_job
        old_repeat_product_card_label_pdf = webapp._repeat_product_card_label_pdf
        calls = []
        print_job = object()
        webapp.create_all = lambda: None
        webapp.load_config = lambda: AppConfig(secret_key="test-secret")
        webapp.session_scope = lambda: FakeSessionScope()
        webapp.get_or_create_settings = lambda _session, _user_id: object()
        webapp.settings_to_app_config = lambda _settings: AppConfig(wb_api_token="token")
        webapp.teksher_mapping_service = FakeTeksherMappingService()
        webapp._get_user_label_print_job = lambda session, print_id, user_id, wb_article: calls.append(
            {
                "print_id": print_id,
                "user_id": user_id,
                "wb_article": wb_article,
            }
        ) or print_job
        webapp._repeat_product_card_label_pdf = lambda request, session, user_id, product_card, loaded_print_job, template: calls.append(
            {
                "repeat_user_id": user_id,
                "product_card": product_card,
                "print_job": loaded_print_job,
                "template": template,
            }
        ) or {
            "ok": True,
            "download_url": "http://testserver/api/labels/pdf/file",
            "labels_count": 1,
            "history": [],
        }
        webapp.product_card_service = FakeProductCardService(
            ProductCardTemplate(
                wb_article="847012873",
                image_url="",
                api_status="",
                wb_summary=WbProductSummary(
                    name="Sport suit",
                    seller_category="Sport suits",
                    wb_article="847012873",
                    tnved="6112120000",
                    country="KG",
                    seller_article="cv_nk_blue_smr",
                    color="blue",
                    composition="polyester 100%",
                    gender="boys",
                    brand="ErLine",
                ),
                rows=[
                    ProductCardMappingRow(
                        barcode="2049271462634",
                        wb_size="38",
                        ru_size="134",
                        teksher_size="",
                        product_type="",
                        gtin="04709055620664",
                        tnved="",
                        country="",
                        vendor_article="",
                        color="",
                        composition="",
                        target_gender="",
                        trademark="",
                        ready_to_mark=0,
                        print_count=0,
                        order_count=0,
                    )
                ],
            )
        )
        try:
            app = webapp.create_app()
            client = TestClient(app)
            client.cookies.set("wb_marks_session", _session_cookie({"user_id": "user-1", "login": "tester"}, "test-secret"))

            response = client.post(
                "/api/product-cards/847012873/label-prints/print-1/repeat",
                json={"template": "Simple"},
            )

            self.assertEqual(200, response.status_code)
            payload = response.json()
            self.assertTrue(payload["ok"])
            self.assertEqual("print-1", calls[0]["print_id"])
            self.assertEqual("847012873", calls[0]["wb_article"])
            self.assertEqual(print_job, calls[1]["print_job"])
            self.assertEqual("simple", calls[1]["template"])
            client.close()
        finally:
            webapp.create_all = old_create_all
            webapp.load_config = old_load_config
            webapp.session_scope = old_session_scope
            webapp.get_or_create_settings = old_get_or_create_settings
            webapp.settings_to_app_config = old_settings_to_app_config
            webapp.product_card_service = old_product_card_service
            webapp.teksher_mapping_service = old_teksher_mapping_service
            webapp._get_user_label_print_job = old_get_user_label_print_job
            webapp._repeat_product_card_label_pdf = old_repeat_product_card_label_pdf

    def test_ready_to_print_counts_use_latest_marking_operation(self) -> None:
        import wb_marks_app.webapp as webapp
        from sqlalchemy import create_engine
        from sqlalchemy.orm import Session
        from wb_marks_app.db import Base
        from wb_marks_app.server_models import LabelPrintJobModel, MarkCodeModel, TeksherOperationModel, WorkflowRunItemModel, WorkflowRunModel
        from wb_marks_app.services.labels import GS

        engine = create_engine("sqlite:///:memory:", future=True)
        Base.metadata.create_all(engine)
        with Session(engine) as session:
            old_run = WorkflowRunModel(
                user_id="user-1",
                draft_id="old",
                source_url="/product-cards/847012873",
                status="completed",
                created_at=datetime(2026, 6, 23, 9, 0, tzinfo=timezone.utc),
            )
            old_item = WorkflowRunItemModel(
                run=old_run,
                barcode="2049271462689",
                vendor_code="cv_nk_blue_smr",
                size="38",
                gtin="04709055620664",
                quantity=3,
                status="completed",
                created_at=datetime(2026, 6, 23, 9, 0, tzinfo=timezone.utc),
            )
            old_marking = TeksherOperationModel(run_item=old_item, operation_kind="marking", status="ACCEPTED")
            failed_run = WorkflowRunModel(
                user_id="user-1",
                draft_id="new",
                source_url="/product-cards/847012873",
                status="partial_failed",
                created_at=datetime(2026, 6, 23, 10, 0, tzinfo=timezone.utc),
            )
            failed_item = WorkflowRunItemModel(
                run=failed_run,
                barcode="2049271462689",
                vendor_code="cv_nk_blue_smr",
                size="38",
                gtin="04709055620664",
                quantity=5,
                status="failed",
                created_at=datetime(2026, 6, 23, 10, 0, tzinfo=timezone.utc),
            )
            failed_marking = TeksherOperationModel(run_item=failed_item, operation_kind="marking", status="REJECTED")
            accepted_run = WorkflowRunModel(
                user_id="user-1",
                draft_id="accepted",
                source_url="/product-cards/847012873",
                status="completed",
                created_at=datetime(2026, 6, 23, 11, 0, tzinfo=timezone.utc),
            )
            accepted_item = WorkflowRunItemModel(
                run=accepted_run,
                barcode="2049271462641",
                vendor_code="cv_nk_blue_smr",
                size="40",
                gtin="04709055620671",
                quantity=2,
                status="completed",
                created_at=datetime(2026, 6, 23, 11, 0, tzinfo=timezone.utc),
            )
            accepted_marking = TeksherOperationModel(run_item=accepted_item, operation_kind="marking", status="ACCEPTED")
            valid_mark_code_1 = "0104709055620664215YudSpca<mc9X" + GS + "91EE12" + GS + "92" + ("A" * 44)
            valid_mark_code_2 = "0104709055620664215YudSpca<mc9Y" + GS + "91EE12" + GS + "92" + ("B" * 44)
            accepted_code_1 = MarkCodeModel(
                run_item=accepted_item,
                position=1,
                mark_code=f"1;{valid_mark_code_1}",
            )
            accepted_code_2 = MarkCodeModel(
                run_item=accepted_item,
                position=2,
                mark_code=f"2;{valid_mark_code_2}",
            )
            session.add_all([old_marking, failed_marking, accepted_marking, accepted_code_1, accepted_code_2])
            session.commit()

            counts = webapp._product_card_ready_to_print_counts(session, "user-1", "847012873")
            print_job = LabelPrintJobModel(
                user_id="user-1",
                wb_article="847012873",
                wb_size="40",
                gtin="04709055620671",
                barcode="2049271462641",
                vendor_code="cv_nk_blue_smr",
                quantity=2,
                template="srad",
                file_id="f" * 32,
                file_name="labels.pdf",
                status="created",
                created_at=datetime(2026, 6, 23, 12, 0, tzinfo=timezone.utc),
            )
            session.add(print_job)
            session.commit()
            counts_after_print = webapp._product_card_ready_to_print_counts(session, "user-1", "847012873")

        self.assertEqual(0, counts["38"])
        self.assertEqual(2, counts["40"])
        self.assertEqual(0, counts_after_print["40"])

    def test_product_card_label_pdf_helper_saves_print_history(self) -> None:
        import wb_marks_app.webapp as webapp
        from pathlib import Path
        from sqlalchemy import create_engine
        from sqlalchemy.orm import Session
        from wb_marks_app.db import Base
        from wb_marks_app.server_models import AppSettingsModel, LabelPrintJobModel, MarkCodeModel, TeksherOperationModel, WorkflowRunItemModel, WorkflowRunModel
        from wb_marks_app.services.labels import GS
        from wb_marks_app.services.product_cards import ProductCardTemplate, ProductCardMappingRow, WbProductSummary

        class FakeRequest:
            def url_for(self, _name, file_id):
                return f"http://testserver/api/labels/pdf/{file_id}"

        valid_mark_code = "0104709055620664215YudSpca<mc9X" + GS + "91EE12" + GS + "92" + ("A" * 44)
        product_card = ProductCardTemplate(
            wb_article="847012873",
            image_url="",
            api_status="",
            wb_summary=WbProductSummary(
                name="Sport suit",
                seller_category="Sport suits",
                wb_article="847012873",
                tnved="6112120000",
                country="KG",
                seller_article="cv_nk_blue_smr",
                color="blue",
                composition="polyester 100%",
                gender="boys",
                brand="ErLine",
            ),
            rows=[
                ProductCardMappingRow(
                    barcode="2049271462634",
                    wb_size="38",
                    ru_size="134",
                    teksher_size="",
                    product_type="",
                    gtin="04709055620664",
                    tnved="6112120000",
                    country="KG",
                    vendor_article="cv_nk_blue_smr",
                    color="ГОЛУБОЙ",
                    composition="polyester 100%",
                    target_gender="",
                    trademark="ErLine",
                    ready_to_mark=0,
                    print_count=1,
                    order_count=0,
                )
            ],
        )
        old_render_labels_pdf = webapp.render_labels_pdf
        old_label_pdf_file_path = webapp._label_pdf_file_path
        engine = create_engine("sqlite:///:memory:", future=True)
        Base.metadata.create_all(engine)
        with TemporaryDirectory() as tmp:
            captured_labels = []

            def fake_render_labels_pdf(labels, template):
                captured_labels.extend(labels)
                return b"%PDF\n/Type /Page\n/Type /Pages\n"

            webapp.render_labels_pdf = fake_render_labels_pdf
            webapp._label_pdf_file_path = lambda _user_id, file_id: Path(tmp) / f"{file_id}.pdf"
            try:
                with Session(engine) as session:
                    settings = AppSettingsModel(
                        user_id="user-1",
                        supplier_name='ОсОО "ЭмирЛайн"',
                        production_address="КР, г. Бишкек, ул. Тестовая, 1",
                    )
                    run = WorkflowRunModel(
                        user_id="user-1",
                        draft_id="run-1",
                        source_url="/product-cards/847012873",
                        status="completed",
                        created_at=datetime(2026, 6, 23, 10, 0, tzinfo=timezone.utc),
                    )
                    item = WorkflowRunItemModel(
                        run=run,
                        barcode="2049271462634",
                        vendor_code="cv_nk_blue_smr",
                        size="38",
                        gtin="04709055620664",
                        quantity=1,
                        status="completed",
                        wb_item_name="Sport suit",
                        created_at=datetime(2026, 6, 23, 10, 0, tzinfo=timezone.utc),
                    )
                    operation = TeksherOperationModel(run_item=item, operation_kind="marking", status="ACCEPTED")
                    mark_code = MarkCodeModel(run_item=item, position=1, mark_code=valid_mark_code)
                    session.add_all([settings, operation, mark_code])
                    session.commit()

                    result = webapp._create_product_card_label_pdf(FakeRequest(), session, "user-1", product_card, "srad")

                    self.assertTrue(result["ok"])
                    self.assertEqual(1, result["labels_count"])
                    self.assertEqual(1, result["pages_count"])
                    self.assertTrue((Path(tmp) / f"{result['file_id']}.pdf").exists())
                    jobs = session.query(LabelPrintJobModel).all()
                    self.assertEqual(1, len(jobs))
                    self.assertEqual("38", jobs[0].wb_size)
                    self.assertEqual("04709055620664", jobs[0].gtin)
                    self.assertEqual(1, jobs[0].quantity)
                    self.assertEqual("srad", jobs[0].template)
                    self.assertEqual("SRad", result["history"][0]["template"])
                    self.assertEqual(1, len(captured_labels))
                    self.assertEqual("Sport suits", captured_labels[0].item_name)
                    self.assertEqual("голубой", captured_labels[0].color)
                    self.assertEqual('ОсОО "ЭмирЛайн"', captured_labels[0].supplier_name)
                    self.assertEqual(date.today().strftime("%d.%m.%Y"), captured_labels[0].production_date)
                    self.assertEqual("КР, г. Бишкек, ул. Тестовая, 1", captured_labels[0].supplier_address)
            finally:
                webapp.render_labels_pdf = old_render_labels_pdf
                webapp._label_pdf_file_path = old_label_pdf_file_path


class FakeSessionScope:
    def __enter__(self):
        return object()

    def __exit__(self, exc_type, exc, traceback):
        return False


class FakeWorkflowRunService:
    def __init__(self) -> None:
        self.calls = []

    def create_product_card_run(self, user_id, wb_article, product_card, rows):
        self.calls.append(
            {
                "user_id": user_id,
                "wb_article": wb_article,
                "product_card": product_card,
                "rows": rows,
            }
        )
        return "run-1"


class FakeProductCardService:
    def __init__(self, product_card) -> None:
        self.product_card = product_card
        self.calls = []

    def build_template(self, wb_article, config):
        self.calls.append((wb_article, config))
        return self.product_card


class FakeTeksherMappingService:
    def __init__(self, version: int = 1, latest_rows=None) -> None:
        self.version = version
        self.latest_rows = latest_rows or []
        self.saved = []
        self.applied = []

    def apply_latest(self, session, user_id, product_card):
        return product_card

    def latest_payload(self, session, user_id, wb_article):
        if not self.latest_rows:
            return 0, []
        return self.version, self.latest_rows

    def apply_payload(self, product_card, saved_rows, version, teksher_rows_by_gtin=None):
        self.applied.append(
            {
                "product_card": product_card,
                "saved_rows": saved_rows,
                "version": version,
                "teksher_rows_by_gtin": teksher_rows_by_gtin,
            }
        )
        if not saved_rows:
            return product_card
        rows = []
        for row in product_card.rows:
            saved = next((item for item in saved_rows if item.get("wb_size") == row.wb_size), None)
            if saved is None:
                rows.append(row)
                continue
            source = (teksher_rows_by_gtin or {}).get(saved.get("gtin"), {})
            rows.append(
                replace(
                    row,
                    full_name=source.get("full_name", "") or saved.get("full_name", "") or _full_name_from_product_type(
                        saved.get("product_type", "")
                    ),
                    teksher_size=source.get("teksher_size", ""),
                    product_type=source.get("product_type", ""),
                    gtin=source.get("gtin", "") or saved.get("gtin", ""),
                    tnved=source.get("tnved", ""),
                    country=source.get("country", ""),
                    vendor_article=source.get("vendor_article", ""),
                    color=source.get("color", ""),
                    composition=source.get("composition", ""),
                    target_gender=source.get("target_gender", ""),
                    trademark=source.get("trademark", ""),
                )
            )
        return replace(product_card, rows=rows, has_teksher_mapping=True, mapping_version=version)

    def save_version(self, session, user_id, product_card, rows_payload, source="product_card"):
        self.saved.append(
            {
                "session": session,
                "user_id": user_id,
                "product_card": product_card,
                "rows": rows_payload,
                "source": source,
            }
        )
        return self.version, len(rows_payload)


class FakeTeksherProductService:
    def __init__(self, draft_ids=None, existing: bool = False, mapping_rows_by_gtin=None, existing_gtins=None) -> None:
        self.draft_ids = draft_ids or []
        self.existing = existing
        self.mapping_rows_by_gtin = mapping_rows_by_gtin or {}
        self.existing_gtins = set(existing_gtins or [])
        self.calls = []
        self.gtins = []

    def product_mapping_rows_by_gtins(self, gtins, config):
        self.gtins = list(gtins)
        return self.mapping_rows_by_gtin

    def product_draft_preview_for_mapping(self, product_card, rows_payload, config):
        rows = list(rows_payload)
        existing_gtins = []
        create_gtins = []
        for row in rows:
            gtin = str(row.get("gtin") or "")
            if self.existing or gtin in self.existing_gtins:
                existing_gtins.append(gtin)
            else:
                create_gtins.append(gtin)
        return {
            "existing_gtins": existing_gtins,
            "create_gtins": create_gtins,
            "draft_fields": {
                "full_name": "Sport suit",
                "tnved": "6112120000",
                "country": "КЫРГЫЗСТАН",
                "manufacturer_inn": "12345678901234",
                "manufacturer_full_name": "ОсОО ЭрЛайн",
                "trademark": "ErLine",
                "product_type": "КОСТЮМ СПОРТИВНЫЙ",
                "vendor_article": "cv_nk_white_smr",
                "regulation": "ТР ТС 017/2011",
                "size_unit": "МЕЖДУНАРОДНЫЙ",
                "composition": "polyester 100%",
                "color": "БЕЛЫЙ",
                "target_gender": "МУЖСКОЙ",
            },
            "dictionaries": {
                "country": [{"value": "КЫРГЫЗСТАН", "label": "KG"}],
                "target_gender": [{"value": "МУЖСКОЙ", "label": ""}],
            },
        }

    def rows_with_product_draft_fields(self, rows_payload, draft_fields):
        rows = [dict(row) for row in rows_payload]
        if not isinstance(draft_fields, dict):
            return rows
        for row in rows:
            for key in (
                "full_name",
                "tnved",
                "country",
                "product_type",
                "vendor_article",
                "color",
                "composition",
                "target_gender",
                "trademark",
            ):
                if draft_fields.get(key):
                    row[key] = draft_fields[key]
        return rows

    def ensure_product_drafts_result_for_mapping(self, product_card, rows_payload, config):
        draft_ids = self.ensure_product_drafts_for_mapping(product_card, rows_payload, config)
        if self.existing:
            return {
                "draft_ids": [],
                "created_gtins": [],
                "existing_gtins": [str(row.get("gtin") or "") for row in rows_payload],
            }
        if self.existing_gtins:
            return {
                "draft_ids": draft_ids,
                "created_gtins": [str(row.get("gtin") or "") for row in rows_payload if str(row.get("gtin") or "") not in self.existing_gtins],
                "existing_gtins": [str(row.get("gtin") or "") for row in rows_payload if str(row.get("gtin") or "") in self.existing_gtins],
            }
        return {
            "draft_ids": draft_ids,
            "created_gtins": [str(row.get("gtin") or "") for row in rows_payload],
            "existing_gtins": [],
        }

    def ensure_product_drafts_for_mapping(self, product_card, rows_payload, config):
        self.calls.append(
            {
                "product_card": product_card,
                "rows": rows_payload,
                "config": config,
            }
        )
        if self.existing:
            return []
        return self.draft_ids


def _build_gtin_upload_workbook(variety: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Перечень продукции"])
    sheet.append(["Наименование предприятия", "ОсОО"])
    sheet.append(["Описание"])
    sheet.append(
        [
            "ТИП",
            "GTIN",
            "БРЕНД",
            "СУБ-БРЕНД",
            "ЯЗЫК",
            "ФУНКЦИОНАЛЬНОЕ НАЗВАНИЕ",
            "РАЗНОВИДНОСТЬ",
        ]
    )
    sheet.append(["Единичная упаковка", "04709055620626", "ErLine", "", "Русский", "Костюмы спортивные", variety])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _full_name_from_product_type(value) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    if text.casefold() in {"костюм спортивный", "костюмы спортивные"}:
        return "Костюм спортивный"
    return text[:1].upper() + text[1:]


def _session_cookie(payload: dict[str, str], secret_key: str) -> str:
    data = b64encode(json.dumps(payload).encode("utf-8"))
    return TimestampSigner(secret_key).sign(data).decode("utf-8")


if __name__ == "__main__":
    unittest.main()
